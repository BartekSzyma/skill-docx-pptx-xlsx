"""Prezentacja o spółce, zbudowana wyłącznie z ledgera i wyniku wyceny.

Użycie: python narzedzia/generuj_pptx.py TICKER

Kroki:
  1. sprawdza weryfikację ledgera i to, że wycena.json powstała z tego samego ledgera,
  2. zapisuje TICKER/wyniki/TICKER_prezentacja.pptx (16:9); przy każdej liczbie i fakcie
     drobnym tekstem id wpisu ledgera, np. [F-07], albo [wycena] dla wyników z wycena.json,
  3. robi podgląd slajdów do PNG w TICKER/wyniki/podglad/ (slajd-NN.png).

Liczby o spółce pochodzą tylko z ledger.json i wyniki/wycena.json.
"""

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from bramka import KATALOG_GLOWNY, wymagaj_weryfikacji
from formaty import NBSP, liczba, wartosc_z_jedn
from libreoffice import podglad_png

# Wygląd
CIEMNY = RGBColor(0x1E, 0x2A, 0x36)      # grafit: tło slajdów tytułowych, tytuły
AKCENT = RGBColor(0x0F, 0x7C, 0x80)      # morski: wyróżnienia, słupki
AKCENT_JASNY = RGBColor(0x8F, 0xC9, 0xC6)
TLO_KARTY = RGBColor(0xF1, 0xF4, 0xF6)
TEKST = RGBColor(0x26, 0x32, 0x3C)
SZARY = RGBColor(0x6B, 0x77, 0x85)
CZERWONY = RGBColor(0xB0, 0x3A, 0x2E)
BIALY = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
SZER, WYS = Inches(13.333), Inches(7.5)
MARGINES = Inches(0.6)


# ---------------------------------------------------------------- dane


class Dane:
    def __init__(self, ticker):
        self.ticker = ticker
        katalog = KATALOG_GLOWNY / ticker
        self.wpisy = json.loads((katalog / "ledger.json").read_text(encoding="utf-8"))
        self.po_kluczu = {w["klucz"]: w for w in self.wpisy}
        self.po_id = {w["id"]: w for w in self.wpisy}
        plik_wyceny = katalog / "wyniki" / "wycena.json"
        if not plik_wyceny.exists():
            sys.exit(f"STOP: brak {plik_wyceny}. Najpierw uruchom: python narzedzia/generuj_xlsx.py {ticker}")
        self.wycena = json.loads(plik_wyceny.read_text(encoding="utf-8"))
        status = json.loads((katalog / "weryfikacja" / "status.json").read_text(encoding="utf-8"))
        if self.wycena.get("sha256_ledgera") != status["sha256_ledgera"]:
            sys.exit("STOP: wycena.json powstała z innej wersji ledgera. Uruchom ponownie: "
                     f"python narzedzia/generuj_xlsx.py {ticker}")

    def __getitem__(self, klucz):
        if klucz not in self.po_kluczu:
            sys.exit(f"STOP: w ledgerze {self.ticker} brakuje wpisu o kluczu '{klucz}'.")
        return self.po_kluczu[klucz]

    def grafiki(self):
        return [w for w in self.wpisy if w["klasa"] == "grafika"]


# ---------------------------------------------------------------- elementy slajdów


def tekst(slajd, x, y, w, h, tresc, rozmiar=16, kolor=TEKST, pogrub=False, wyrownanie=PP_ALIGN.LEFT,
          kotwica=MSO_ANCHOR.TOP, kursywa=False):
    """Pole tekstowe; tresc to napis albo lista akapitów, akapit to napis albo lista fragmentów (tekst, opcje)."""
    pole = slajd.shapes.add_textbox(x, y, w, h)
    tf = pole.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = kotwica
    akapity = tresc if isinstance(tresc, list) else [tresc]
    for i, akapit in enumerate(akapity):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = wyrownanie
        fragmenty = akapit if isinstance(akapit, list) else [(akapit, {})]
        for t, opcje in fragmenty:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opcje.get("rozmiar", rozmiar))
            f.bold = opcje.get("pogrub", pogrub)
            f.italic = opcje.get("kursywa", kursywa)
            f.color.rgb = opcje.get("kolor", kolor)
        pierwsze = akapit[0][1] if isinstance(akapit, list) else {}
        if "odstep" in pierwsze:
            p.space_after = Pt(pierwsze["odstep"])
        if pierwsze.get("punkt"):
            ustaw_punktor(p, pierwsze.get("rozmiar", rozmiar))
    return pole


def ustaw_punktor(p, rozmiar):
    """Punktor akapitu (kwadrat w kolorze akcentu) z wcięciem wiszącym: kolejne linie zaczynają się pod tekstem."""
    pPr = p._p.get_or_add_pPr()
    wciecie = int(Pt(rozmiar) * 1.4)
    pPr.set("marL", str(wciecie))
    pPr.set("indent", str(-wciecie))
    for znacznik, atrybuty in (("a:buClr", None), ("a:buSzPct", {"val": "75000"}), ("a:buFont", {"typeface": "Arial"}),
                               ("a:buChar", {"char": "■"})):
        el = pPr.makeelement(qn(znacznik), atrybuty or {})
        if znacznik == "a:buClr":
            kolor = el.makeelement(qn("a:srgbClr"), {"val": str(AKCENT)})
            el.append(kolor)
        pPr.append(el)


def id_tag(*ids):
    return (f"  [{', '.join(ids)}]", {"rozmiar": 10, "kolor": SZARY, "pogrub": False})


def punkt(tresc, *ids, rozmiar=16, odstep=10):
    """Punkt listy: kwadracik w kolorze akcentu, treść, id wpisów drobnym tekstem."""
    return [(tresc, {"rozmiar": rozmiar, "punkt": True, "odstep": odstep}), id_tag(*ids)]


def prostokat(slajd, x, y, w, h, kolor, zaokraglony=True):
    ksztalt = slajd.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if zaokraglony else MSO_SHAPE.RECTANGLE, x, y, w, h)
    ksztalt.fill.solid()
    ksztalt.fill.fore_color.rgb = kolor
    ksztalt.line.fill.background()
    ksztalt.shadow.inherit = False
    if zaokraglony:
        ksztalt.adjustments[0] = 0.08
    return ksztalt


def tlo(slajd, kolor):
    slajd.background.fill.solid()
    slajd.background.fill.fore_color.rgb = kolor


def tytul(slajd, napis, podtytul=None):
    tekst(slajd, MARGINES, Inches(0.45), SZER - 2 * MARGINES, Inches(0.75), napis, rozmiar=34, kolor=CIEMNY, pogrub=True)
    if podtytul:
        tekst(slajd, MARGINES, Inches(1.15), SZER - 2 * MARGINES, Inches(0.4), podtytul, rozmiar=15, kolor=SZARY)


def stopka(slajd, numer, ticker, ciemne=False):
    kolor = AKCENT_JASNY if ciemne else SZARY
    tekst(slajd, MARGINES, WYS - Inches(0.45), Inches(8), Inches(0.3),
          f"{ticker} - pakiet analityczny. Dane: ledger.json i wycena.json; [id] = wpis ledgera.", rozmiar=10, kolor=kolor)
    tekst(slajd, SZER - MARGINES - Inches(1), WYS - Inches(0.45), Inches(1), Inches(0.3), str(numer), rozmiar=10,
          kolor=kolor, wyrownanie=PP_ALIGN.RIGHT)


def karta_liczby(slajd, x, y, w, h, wartosc, opis, ids):
    prostokat(slajd, x, y, w, h, TLO_KARTY)
    tekst(slajd, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.65), wartosc, rozmiar=30, kolor=AKCENT, pogrub=True)
    tekst(slajd, x + Inches(0.25), y + Inches(0.85), w - Inches(0.5), h - Inches(0.95),
          [[(opis, {"rozmiar": 14}), id_tag(*ids)]], kolor=TEKST)


def styl_wykresu(wykres, legenda=True, rozmiar=12):
    wykres.font.name = FONT
    wykres.font.size = Pt(rozmiar)
    wykres.font.color.rgb = TEKST
    wykres.has_legend = legenda
    if legenda:
        wykres.legend.position = XL_LEGEND_POSITION.BOTTOM
        wykres.legend.include_in_layout = False
        wykres.legend.font.size = Pt(rozmiar)
    os_w = wykres.value_axis
    os_w.has_major_gridlines = True
    os_w.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE6)
    os_w.format.line.fill.background()
    os_w.tick_labels.font.size = Pt(rozmiar - 1)
    os_w.tick_labels.font.color.rgb = SZARY
    os_k = wykres.category_axis
    os_k.format.line.color.rgb = RGBColor(0xBF, 0xC6, 0xCC)
    os_k.tick_labels.font.size = Pt(rozmiar)


def etykiety(seria, format_liczb='#,##0;-#,##0', pozycja=XL_LABEL_POSITION.OUTSIDE_END, rozmiar=11):
    seria.data_labels.show_value = True
    seria.data_labels.number_format = format_liczb
    seria.data_labels.number_format_is_linked = False
    seria.data_labels.position = pozycja
    seria.data_labels.font.size = Pt(rozmiar)
    seria.data_labels.font.color.rgb = TEKST


def zdjecie(slajd, plik, x, y, w_max, h_max):
    """Wstawia zdjęcie dopasowane do ramki (bez zniekształceń) i zwraca jego położenie."""
    with Image.open(plik) as obraz:
        szer, wys = obraz.size
    skala = min(w_max / szer, h_max / wys)
    w, h = int(szer * skala), int(wys * skala)
    x0, y0 = x + (w_max - w) // 2, y
    slajd.shapes.add_picture(str(plik), x0, y0, w, h)
    return x0, y0, w, h


# ---------------------------------------------------------------- slajdy


def slajd_tytulowy(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tlo(s, CIEMNY)
    dokument = D["revenue_2025"]["zrodlo"]["dokument"]
    spolka = dokument.split(" - ")[0]
    tekst(s, MARGINES, Inches(2.0), Inches(11), Inches(1.0), spolka, rozmiar=48, kolor=BIALY, pogrub=True)
    tekst(s, MARGINES, Inches(3.0), Inches(11), Inches(0.6), f"Pakiet analityczny  |  {D.ticker}", rozmiar=26, kolor=AKCENT_JASNY)
    r01 = D["share_price"]
    tekst(s, MARGINES, Inches(4.3), Inches(11.5), Inches(1.2), [
        [(f"Dane finansowe: {dokument}", {"rozmiar": 16, "kolor": BIALY}), id_tag(D["revenue_2025"]["id"])],
        [(f"Dane rynkowe: kurs z dnia {r01['okres']}, pobrane {r01['data_pobrania']}", {"rozmiar": 16, "kolor": BIALY}),
         id_tag(r01["id"])],
        [(f"Wycena DCF z dnia {D.wycena['data_wyceny']}", {"rozmiar": 16, "kolor": BIALY}), id_tag("wycena")],
    ])
    stopka(s, nr, D.ticker, ciemne=True)


def slajd_liczby(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, "Spółka w liczbach", "Przychody i zysk operacyjny w latach 2023-2025, mln USD")
    lata = ["2023", "2024", "2025"]
    rev = [D[f"revenue_{r}"] for r in lata]
    opi = [D[f"operating_income_{r}"] for r in lata]
    cd = CategoryChartData()
    cd.categories = lata
    cd.add_series("Przychody", [w["wartosc"] for w in rev])
    cd.add_series("Zysk operacyjny", [w["wartosc"] for w in opi])
    ramka = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGINES, Inches(1.75), Inches(7.6), Inches(4.65), cd)
    wykres = ramka.chart
    styl_wykresu(wykres)
    wykres.value_axis.tick_labels.number_format = '#,##0;-#,##0'
    wykres.value_axis.tick_labels.number_format_is_linked = False
    wykres.plots[0].gap_width = 60
    wykres.plots[0].overlap = -10
    for seria, kolor in zip(wykres.series, (AKCENT, CIEMNY)):
        seria.format.fill.solid()
        seria.format.fill.fore_color.rgb = kolor
        etykiety(seria)
    tekst(s, MARGINES, Inches(6.45), Inches(7.6), Inches(0.3),
          [[("Źródło: rachunek zysków i strat, raport 10-K, s. 89", {"rozmiar": 10, "kolor": SZARY}),
            id_tag(*[w["id"] for w in rev + opi])]])

    x, w = Inches(8.7), Inches(4.0)
    ni, marza, got = D["net_income_2025"], D["operating_margin_2025"], D["cash_and_marketable_securities_2025"]
    karta_liczby(s, x, Inches(1.75), w, Inches(1.4), wartosc_z_jedn(ni), f"zysk netto w {ni['okres']} r.", [ni["id"]])
    karta_liczby(s, x, Inches(3.35), w, Inches(1.4), wartosc_z_jedn(marza), f"marża operacyjna w {marza['okres']} r.", [marza["id"]])
    karta_liczby(s, x, Inches(4.95), w, Inches(1.4), wartosc_z_jedn(got),
                 "gotówka, ekwiwalenty i papiery wartościowe na koniec 2025 r.", [got["id"]])
    stopka(s, nr, D.ticker)


def slajd_ai(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, "AI i aplikacje: źródło zysków", "Family of Apps zarabia, Reality Labs inwestuje")
    si, dap, foa_r, foa_o, rl_r, rl_o = (D[k] for k in ("superintelligence_vision", "family_daily_active_people_dec_2025",
                                                         "foa_revenue_2025", "foa_operating_income_2025", "rl_revenue_2025",
                                                         "rl_operating_income_2025"))
    tekst(s, MARGINES, Inches(1.85), Inches(6.0), Inches(4.3), [
        punkt("Cel spółki w AI: budowa superinteligencji, czyli AI przewyższającej ludzką inteligencję", si["id"]),
        punkt(f"{liczba(dap['wartosc'], 2)} mld osób dziennie korzysta z aplikacji Meta (średnio w grudniu 2025 r.)", dap["id"]),
        punkt(f"Przychody segmentu Family of Apps w 2025 r.: {wartosc_z_jedn(foa_r)}", foa_r["id"]),
        punkt(f"Zysk operacyjny Family of Apps: {wartosc_z_jedn(foa_o)}; strata Reality Labs: {wartosc_z_jedn(rl_o)}",
              foa_o["id"], rl_o["id"]),
    ])
    cd = CategoryChartData()
    cd.categories = ["Family of Apps", "Reality Labs"]
    # Przychody (wartości dodatnie): LibreOffice rysuje ujemne słupki z python-pptx jako dodatnie, więc strata
    # Reality Labs jest w punktach po lewej, a nie na wykresie.
    cd.add_series("Przychody 2025, mln USD", [foa_r["wartosc"], rl_r["wartosc"]])
    ramka = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.2), Inches(1.75), Inches(5.5), Inches(4.65), cd)
    wykres = ramka.chart
    styl_wykresu(wykres, legenda=False)
    wykres.has_title = True
    wykres.chart_title.text_frame.text = "Przychody segmentów w 2025 r., mln USD"
    wykres.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    wykres.chart_title.text_frame.paragraphs[0].runs[0].font.bold = False
    wykres.value_axis.tick_labels.number_format = '#,##0;-#,##0'
    wykres.value_axis.tick_labels.number_format_is_linked = False
    seria = wykres.series[0]
    seria.invert_if_negative = False
    for i, kolor in enumerate((AKCENT, CZERWONY)):
        pt = seria.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = kolor
    etykiety(seria)
    tekst(s, Inches(7.2), Inches(6.45), Inches(5.5), Inches(0.3),
          [[("Źródło: nota 15, raport 10-K, s. 126", {"rozmiar": 10, "kolor": SZARY}), id_tag(foa_r["id"], rl_r["id"])]])
    stopka(s, nr, D.ticker)


def slajd_projekt(prs, D, nr, g, napis, podtytul, punkty):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, napis, podtytul)
    plik = KATALOG_GLOWNY / g["plik"]
    x_zdj, y_zdj = MARGINES, Inches(1.75)
    _, _, _, h = zdjecie(s, plik, x_zdj, y_zdj, Inches(6.4), Inches(4.3))
    tekst(s, x_zdj, y_zdj + h + Inches(0.08), Inches(6.4), Inches(0.5), [
        [(f"{g['opis']}. Fot. {g['autor']}, licencja {g['licencja']}, Wikimedia Commons", {"rozmiar": 10, "kolor": SZARY}),
         id_tag(g["id"])]])
    prostokat(s, Inches(7.35), Inches(1.75), Inches(5.4), Inches(4.65), TLO_KARTY)
    tekst(s, Inches(7.65), Inches(2.05), Inches(4.8), Inches(4.1), [punkt(t, *ids, rozmiar=17, odstep=14) for t, ids in punkty])
    stopka(s, nr, D.ticker)


def slajd_wycena(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W = D.wycena
    tytul(s, "Wycena DCF: wartość wobec kursu", "Wynik modelu z pliku META_DCF.xlsx; zależy od przyjętych założeń")
    cd = CategoryChartData()
    cd.categories = ["Wartość z DCF", "Kurs rynkowy"]
    cd.add_series("USD za akcję", [W["wartosc_jednej_akcji_usd"], W["kurs_usd"]])
    ramka = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGINES, Inches(1.75), Inches(6.2), Inches(2.6), cd)
    wykres = ramka.chart
    styl_wykresu(wykres, legenda=False)
    wykres.value_axis.visible = False
    wykres.value_axis.has_major_gridlines = False
    wykres.value_axis.minimum_scale = 0
    wykres.plots[0].gap_width = 45
    wykres.category_axis.reverse_order = True  # wartość z DCF na górze, jak w kolejności danych
    seria = wykres.series[0]
    for i, kolor in enumerate((AKCENT, CIEMNY)):
        seria.points[i].format.fill.solid()
        seria.points[i].format.fill.fore_color.rgb = kolor
    etykiety(seria, '#,##0.00" USD"', rozmiar=14)
    tekst(s, MARGINES, Inches(4.45), Inches(6.2), Inches(1.1), [
        [(f"{liczba(W['roznica_proc'], 1)}%", {"rozmiar": 36, "kolor": CZERWONY if W["roznica_proc"] < 0 else AKCENT, "pogrub": True}),
         ("  różnica wartości z DCF wobec kursu", {"rozmiar": 16}), id_tag("wycena", D["share_price"]["id"])],
    ])
    tekst(s, MARGINES, Inches(5.4), Inches(6.2), Inches(1.0), [
        [(f"Wartość firmy {liczba(W['wartosc_firmy_mln_usd'])} mln USD, kapitał własny "
          f"{liczba(W['wartosc_kapitalu_mln_usd'])} mln USD; {liczba(W['udzial_wartosci_rezydualnej_proc'], 0)}% wartości firmy "
          "pochodzi z okresu po 2030 r.", {"rozmiar": 14}), id_tag("wycena")]])

    x, w = Inches(7.25), Inches(5.5)
    prostokat(s, x, Inches(1.75), w, Inches(4.65), TLO_KARTY)
    tekst(s, x + Inches(0.3), Inches(1.95), w - Inches(0.6), Inches(0.75), [
        [("To są ZAŁOŻENIA, nie fakty", {"rozmiar": 18, "pogrub": True, "kolor": CZERWONY})],
        [("Przyjęte przez nas, z uzasadnieniem w ledgerze; zmiana założeń zmienia wynik.", {"rozmiar": 12, "kolor": SZARY})]])
    z = {k: D[k] for k in ("revenue_growth_2026", "revenue_growth_2030", "operating_margin_terminal",
                           "capex_pct_revenue_2026", "beta", "equity_risk_premium", "terminal_growth")}
    rf = D["us_treasury_10y_yield"]
    pozycje = [
        (f"Wzrost przychodów: {liczba(z['revenue_growth_2026']['wartosc'])}% w 2026 r., {liczba(z['revenue_growth_2030']['wartosc'])}% w 2030 r.",
         [z["revenue_growth_2026"]["id"], z["revenue_growth_2030"]["id"]]),
        (f"Marża operacyjna po 2030 r.: {liczba(z['operating_margin_terminal']['wartosc'])}%", [z["operating_margin_terminal"]["id"]]),
        (f"Nakłady inwestycyjne w 2026 r.: {liczba(z['capex_pct_revenue_2026']['wartosc'])}% przychodów",
         [z["capex_pct_revenue_2026"]["id"]]),
        (f"WACC {liczba(W['wacc_proc'], 1)}% = rentowność obligacji {liczba(rf['wartosc'], 2)}% + beta "
         f"{liczba(z['beta']['wartosc'], 1)} × premia {liczba(z['equity_risk_premium']['wartosc'], 1)}%",
         ["wycena", rf["id"], z["beta"]["id"], z["equity_risk_premium"]["id"]]),
        (f"Wzrost po 2030 r.: {liczba(z['terminal_growth']['wartosc'], 1)}% rocznie", [z["terminal_growth"]["id"]]),
    ]
    tekst(s, x + Inches(0.3), Inches(2.95), w - Inches(0.6), Inches(3.3), [punkt(t, *ids, rozmiar=15, odstep=9) for t, ids in pozycje])
    stopka(s, nr, D.ticker)


def slajd_ryzyka(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, "Ryzyka i źródła")
    ryzyka = [w for w in D.wpisy if w["klasa"] == "fakt" and w["klucz"].startswith("risk_")]
    y = Inches(1.4)
    for w in ryzyka:
        opis = w["opis"].removeprefix("Ryzyko: ")
        prostokat(s, MARGINES, y, Inches(6.6), Inches(1.45), TLO_KARTY)
        tekst(s, MARGINES + Inches(0.25), y + Inches(0.15), Inches(6.1), Inches(1.2), [
            [(w["wartosc"][0].upper() + w["wartosc"][1:], {"rozmiar": 16, "pogrub": True, "kolor": CZERWONY})],
            [(opis[0].upper() + opis[1:], {"rozmiar": 13}), id_tag(w["id"]),
             (f"  10-K, s. {w['zrodlo']['strona']}", {"rozmiar": 10, "kolor": SZARY})]])
        y += Inches(1.6)

    # Źródła: dokumenty i strony użyte w ledgerze
    strony = sorted({w["zrodlo"]["strona"] for w in D.wpisy if w["klasa"] == "fakt"}, key=lambda s_: (isinstance(s_, str), s_))
    strony_tekst = ", ".join("okładka" if isinstance(s_, str) else str(s_) for s_ in strony)
    dok_10k = D["revenue_2025"]["zrodlo"]
    rynek = [w for w in D.wpisy if w["klasa"] == "dana_rynkowa"]
    graf = D.grafiki()
    zrodla = [[(dok_10k["dokument"], {"rozmiar": 14, "pogrub": True, "odstep": 2})],
              [(f"SEC EDGAR, strony: {strony_tekst}", {"rozmiar": 12, "odstep": 10})]]
    for w in rynek:
        zrodla += [[(w["zrodlo"]["dokument"], {"rozmiar": 14, "pogrub": True, "odstep": 2})],
                   [(f"{w['opis']}; stan na {w['okres']}, pobrano {w['data_pobrania']}", {"rozmiar": 12, "odstep": 10}),
                    id_tag(w["id"])]]
    zrodla += [[("Wikimedia Commons - zdjęcia", {"rozmiar": 14, "pogrub": True, "odstep": 2})],
               [("; ".join(f"{g['autor']} ({g['licencja']})" for g in graf), {"rozmiar": 12, "odstep": 10}),
                id_tag(*[g["id"] for g in graf])]]
    zrodla += [[("Ledger i weryfikacja", {"rozmiar": 14, "pogrub": True, "odstep": 2})],
               [(f"{D.ticker}/ledger.json: każdy wpis ze źródłem i cytatem, sprawdzony dwiema bramkami "
                 "(skrypt i niezależny model)", {"rozmiar": 12})]]
    tekst(s, Inches(7.6), Inches(1.4), Inches(5.1), Inches(0.4), "Źródła", rozmiar=18, kolor=CIEMNY, pogrub=True)
    tekst(s, Inches(7.6), Inches(1.9), Inches(5.1), Inches(4.6), zrodla)
    stopka(s, nr, D.ticker)


# ---------------------------------------------------------------- main


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Użycie: python narzedzia/generuj_pptx.py TICKER")
    ticker = sys.argv[1].upper()
    wymagaj_weryfikacji(ticker)
    D = Dane(ticker)

    prs = Presentation()
    prs.slide_width, prs.slide_height = SZER, WYS
    prs.core_properties.title = f"{ticker} - pakiet analityczny"
    prs.core_properties.language = "pl-PL"

    slajd_tytulowy(prs, D, 1)
    slajd_liczby(prs, D, 2)
    slajd_ai(prs, D, 3)

    grafiki = {g["klucz"]: g for g in D.grafiki()}
    fakty = {k: D[k] for k in ("meta_ray_ban_display_launch", "rl_wearables_opex_share_2026", "rl_revenue_2025",
                               "rl_operating_income_2025", "data_center_locations", "capex_2025", "capex_guidance_2026_low",
                               "capex_guidance_2026_high", "ppe_net_2025")}
    f = fakty
    projekty = [
        ("photo_ai_glasses", "Okulary AI", "Urządzenia noszone to główny kierunek Reality Labs", [
            ("W 2025 r. spółka wprowadziła okulary Meta Ray-Ban Display z wyświetlaczem w szkle",
             [f["meta_ray_ban_display_launch"]["id"]]),
            (f"W 2026 r. ok. {liczba(f['rl_wearables_opex_share_2026']['wartosc'])}% kosztów operacyjnych Reality Labs "
             "ma trafić na urządzenia noszone, w tym okulary AI", [f["rl_wearables_opex_share_2026"]["id"]]),
        ]),
        ("photo_vr_headset", "Gogle VR Meta Quest", "Reality Labs: duże nakłady, małe przychody", [
            (f"Przychody Reality Labs w 2025 r.: {wartosc_z_jedn(f['rl_revenue_2025'])}", [f["rl_revenue_2025"]["id"]]),
            (f"Wynik operacyjny Reality Labs w 2025 r.: {wartosc_z_jedn(f['rl_operating_income_2025'])}",
             [f["rl_operating_income_2025"]["id"]]),
            ("Na VR i Horizon ma trafić pozostała część kosztów operacyjnych Reality Labs w 2026 r.",
             [f["rl_wearables_opex_share_2026"]["id"]]),
        ]),
        ("photo_data_center", "Centra danych", "Infrastruktura pod AI to największy wydatek spółki", [
            (f"Spółka ma {liczba(f['data_center_locations']['wartosc'])} własnych lokalizacji centrów danych na świecie",
             [f["data_center_locations"]["id"]]),
            (f"Nakłady na rzeczowe aktywa trwałe w 2025 r.: {liczba(-f['capex_2025']['wartosc'])}{NBSP}mln USD",
             [f["capex_2025"]["id"]]),
            (f"Plan nakładów inwestycyjnych na 2026 r.: od {liczba(f['capex_guidance_2026_low']['wartosc'])} "
             f"do {liczba(f['capex_guidance_2026_high']['wartosc'])}{NBSP}mld USD",
             [f["capex_guidance_2026_low"]["id"], f["capex_guidance_2026_high"]["id"]]),
        ]),
    ]
    nr = 4
    for klucz, napis, podtytul, punkty in projekty:
        if klucz not in grafiki:
            print(f"Uwaga: brak grafiki '{klucz}' w ledgerze - pomijam slajd '{napis}'.")
            continue
        slajd_projekt(prs, D, nr, grafiki[klucz], napis, podtytul, punkty)
        nr += 1
    slajd_wycena(prs, D, nr)
    slajd_ryzyka(prs, D, nr + 1)

    katalog = KATALOG_GLOWNY / ticker / "wyniki"
    plik = katalog / f"{ticker}_prezentacja.pptx"
    prs.save(plik)
    liczba_slajdow = podglad_png(plik, katalog / "podglad", "slajd", dpi=80)
    print(f"Prezentacja: {plik} ({len(prs.slides)} slajdów), podgląd: {liczba_slajdow} plików PNG w {katalog / 'podglad'}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    main()
