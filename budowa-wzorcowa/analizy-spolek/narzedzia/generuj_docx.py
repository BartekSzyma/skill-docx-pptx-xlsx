"""Jednostronicowa notatka dla zarządu, zbudowana wyłącznie z ledgera i wyniku wyceny.

Użycie: python narzedzia/generuj_docx.py TICKER

Kroki:
  1. sprawdza weryfikację ledgera i to, że wycena.json powstała z tego samego ledgera,
  2. zapisuje TICKER/wyniki/TICKER_notatka_zarzad.docx (A4); przy każdej liczbie id wpisu,
  3. zamienia plik na PDF w LibreOffice i sprawdza, że ma dokładnie 1 stronę,
  4. zapisuje podgląd w TICKER/wyniki/podglad/docx-strona-NN.png.
"""

import json
import sys
import tempfile

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from bramka import KATALOG_GLOWNY, wymagaj_weryfikacji
from formaty import NBSP, liczba, wartosc_z_jedn
from libreoffice import konwertuj, podglad_png

FONT = "Calibri"
CIEMNY = RGBColor(0x1E, 0x2A, 0x36)
AKCENT = RGBColor(0x0F, 0x7C, 0x80)
SZARY = RGBColor(0x6B, 0x77, 0x85)
CZERWONY = RGBColor(0xB0, 0x3A, 0x2E)
TLO_NAGLOWKA_TABELI = "1E2A36"
TLO_WIERSZA = "F1F4F6"
ROZMIAR = 10        # tekst główny, pt
ROZMIAR_DROBNY = 9  # źródła i id wpisów; mniej nie schodzimy
ZASTRZEZENIE = "Materiał edukacyjny, nie stanowi rekomendacji inwestycyjnej."


class Dane:
    def __init__(self, ticker):
        self.ticker = ticker
        katalog = KATALOG_GLOWNY / ticker
        self.wpisy = json.loads((katalog / "ledger.json").read_text(encoding="utf-8"))
        self.po_kluczu = {w["klucz"]: w for w in self.wpisy}
        plik_wyceny = katalog / "wyniki" / "wycena.json"
        if not plik_wyceny.exists():
            sys.exit(f"STOP: brak {plik_wyceny}. Najpierw uruchom: python narzedzia/generuj_xlsx.py {ticker}")
        self.wycena = json.loads(plik_wyceny.read_text(encoding="utf-8"))
        status = json.loads((katalog / "weryfikacja" / "status.json").read_text(encoding="utf-8"))
        if self.wycena.get("sha256_ledgera") != status["sha256_ledgera"]:
            sys.exit(f"STOP: wycena.json powstała z innej wersji ledgera. Uruchom ponownie: python narzedzia/generuj_xlsx.py {ticker}")

    def __getitem__(self, klucz):
        if klucz not in self.po_kluczu:
            sys.exit(f"STOP: w ledgerze {self.ticker} brakuje wpisu o kluczu '{klucz}'.")
        return self.po_kluczu[klucz]


# ---------------------------------------------------------------- elementy


def run(p, tekst, rozmiar=ROZMIAR, pogrub=False, kolor=None, kursywa=False):
    r = p.add_run(tekst)
    r.font.name = FONT
    r.font.size = Pt(rozmiar)
    r.font.bold = pogrub
    r.font.italic = kursywa
    if kolor is not None:
        r.font.color.rgb = kolor
    return r


def ids(p, *identyfikatory):
    run(p, f" [{', '.join(identyfikatory)}]", ROZMIAR_DROBNY, kolor=SZARY)


def akapit(doc, odstep_po=3, styl=None):
    p = doc.add_paragraph(style=styl)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(odstep_po)
    p.paragraph_format.line_spacing = 1.05
    return p


def naglowek_sekcji(doc, tekst):
    p = akapit(doc, odstep_po=2)
    p.paragraph_format.space_before = Pt(7)
    run(p, tekst, 11.5, pogrub=True, kolor=AKCENT)
    return p


def punkt(doc, fragmenty):
    """Punkt listy (styl 'List Bullet' z numeracją Worda); fragmenty: lista (tekst, opcje) albo id."""
    p = akapit(doc, odstep_po=2, styl="List Bullet")
    for f in fragmenty:
        if isinstance(f, tuple) and f[0] == "ID":
            ids(p, *f[1])
        else:
            tekst, opcje = f if isinstance(f, tuple) else (f, {})
            run(p, tekst, **opcje)
    return p


def cieniowanie(komorka, kolor):
    tcPr = komorka._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kolor)
    tcPr.append(shd)


def linia_pod(p, kolor="0F7C80"):
    pPr = p._p.get_or_add_pPr()
    granice = OxmlElement("w:pBdr")
    dol = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", kolor)):
        dol.set(qn(k), v)
    granice.append(dol)
    pPr.append(granice)


# ---------------------------------------------------------------- treść


def zbuduj(D):
    W = D.wycena
    f = {k: D[k] for k in (
        "revenue_2025", "revenue_2024", "operating_margin_2025", "cash_and_marketable_securities_2025", "share_price",
        "capex_2025", "capex_guidance_2026_low", "capex_guidance_2026_high", "superintelligence_vision",
        "meta_ray_ban_display_launch", "rl_wearables_opex_share_2026", "data_center_locations",
        "risk_advertising_dependence", "risk_ai_investments", "risk_regulation",
        "family_daily_active_people_dec_2025")}
    spolka = f["revenue_2025"]["zrodlo"]["dokument"].split(" - ")[0]

    doc = Document()
    sekcja = doc.sections[0]
    sekcja.page_width, sekcja.page_height = Cm(21.0), Cm(29.7)
    sekcja.left_margin = sekcja.right_margin = Cm(1.9)
    sekcja.top_margin, sekcja.bottom_margin = Cm(1.5), Cm(1.4)
    styl = doc.styles["Normal"]
    styl.font.name = FONT
    styl.font.size = Pt(ROZMIAR)
    styl.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # Nagłówek
    p = akapit(doc, odstep_po=0)
    run(p, "Notatka dla zarządu", 10, pogrub=True, kolor=AKCENT)
    p = akapit(doc, odstep_po=1)
    run(p, f"{spolka} ({D.ticker})", 20, pogrub=True, kolor=CIEMNY)
    p = akapit(doc, odstep_po=6)
    run(p, f"Data: {W['data_wyceny']}  |  dane finansowe: raport 10-K za {f['revenue_2025']['okres']} r., "
           f"kurs z {f['share_price']['okres']}", 9.5, kolor=SZARY)
    linia_pod(p)

    # Wnioski
    naglowek_sekcji(doc, "Najważniejsze wnioski")
    punkt(doc, [("Silny, rentowny biznes reklamowy: ", {"pogrub": True}),
                f"przychody wzrosły z {wartosc_z_jedn(f['revenue_2024'])} w 2024 r. do {wartosc_z_jedn(f['revenue_2025'])} "
                f"w 2025 r. przy marży operacyjnej {wartosc_z_jedn(f['operating_margin_2025'])}.",
                ("ID", [f["revenue_2024"]["id"], f["revenue_2025"]["id"], f["operating_margin_2025"]["id"]])])
    punkt(doc, [("Rekordowe inwestycje w AI: ", {"pogrub": True}),
                f"nakłady na środki trwałe wyniosły {liczba(-f['capex_2025']['wartosc'])}{NBSP}mln USD w 2025 r., "
                f"a plan na 2026 r. to od {liczba(f['capex_guidance_2026_low']['wartosc'])} do "
                f"{liczba(f['capex_guidance_2026_high']['wartosc'])}{NBSP}mld USD.",
                ("ID", [f["capex_2025"]["id"], f["capex_guidance_2026_low"]["id"], f["capex_guidance_2026_high"]["id"]])])
    punkt(doc, [("Wycena wymaga ostrożności: ", {"pogrub": True}),
                f"przy naszych założeniach model DCF daje {liczba(W['wartosc_jednej_akcji_usd'], 2)}{NBSP}USD na akcję "
                f"wobec kursu {liczba(W['kurs_usd'], 2)}{NBSP}USD ({liczba(W['roznica_proc'], 1)}%); "
                f"{liczba(W['udzial_wartosci_rezydualnej_proc'])}% wartości firmy przypada na lata po 2030 r., "
                "więc wynik silnie zależy od założeń.",
                ("ID", ["wycena", f["share_price"]["id"]])])

    # Tabela
    naglowek_sekcji(doc, "Kluczowe liczby")
    wiersze = [
        ("Przychody", f["revenue_2025"]["okres"], wartosc_z_jedn(f["revenue_2025"]), f["revenue_2025"]["id"]),
        ("Marża operacyjna", f["operating_margin_2025"]["okres"], wartosc_z_jedn(f["operating_margin_2025"]),
         f["operating_margin_2025"]["id"]),
        ("Gotówka, ekwiwalenty i papiery wartościowe", f["cash_and_marketable_securities_2025"]["okres"],
         wartosc_z_jedn(f["cash_and_marketable_securities_2025"]), f["cash_and_marketable_securities_2025"]["id"]),
        ("Wartość jednej akcji z DCF (założenia, nie fakt)", W["data_wyceny"],
         f"{liczba(W['wartosc_jednej_akcji_usd'], 2)}{NBSP}USD", "wycena"),
        ("Kurs akcji", f["share_price"]["okres"], f"{liczba(f['share_price']['wartosc'], 2)}{NBSP}USD", f["share_price"]["id"]),
        ("Różnica: wartość z DCF wobec kursu", W["data_wyceny"], f"{liczba(W['roznica_proc'], 1)}%", "wycena"),
    ]
    tabela = doc.add_table(rows=1 + len(wiersze), cols=4)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    szerokosci = [Cm(8.6), Cm(2.9), Cm(3.6), Cm(2.1)]
    for i, tekst in enumerate(["Pozycja", "Okres / data", "Wartość", "id"]):
        k = tabela.rows[0].cells[i]
        cieniowanie(k, TLO_NAGLOWKA_TABELI)
        k.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(k.paragraphs[0], tekst, 9.5, pogrub=True, kolor=RGBColor(0xFF, 0xFF, 0xFF))
    for r, wiersz in enumerate(wiersze, start=1):
        for c, tekst in enumerate(wiersz):
            k = tabela.rows[r].cells[c]
            if r % 2 == 0:
                cieniowanie(k, TLO_WIERSZA)
            p = k.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            kolor = SZARY if c == 3 else (CZERWONY if c == 2 and tekst.startswith("-") else None)
            run(p, tekst, ROZMIAR_DROBNY if c == 3 else ROZMIAR, pogrub=(c == 2), kolor=kolor)
    # Szerokości kolumn: siatka tabeli (czyta ją LibreOffice) i szerokość każdej komórki (czyta ją Word).
    tabela.autofit = False
    siatka = tabela._tbl.tblGrid
    for kolumna, s_ in zip(siatka.findall(qn("w:gridCol")), szerokosci):
        kolumna.set(qn("w:w"), str(int(s_.twips)))
    for wiersz in tabela.rows:
        for k, s_ in zip(wiersz.cells, szerokosci):
            k.width = s_

    # Projekty
    naglowek_sekcji(doc, "Projekty")
    punkt(doc, [("Sztuczna inteligencja: ", {"pogrub": True}),
                "spółka deklaruje cel zbudowania superinteligencji, czyli AI przewyższającej ludzką inteligencję, "
                f"a z jej aplikacji korzysta dziennie {liczba(f['family_daily_active_people_dec_2025']['wartosc'], 2)}{NBSP}mld osób.",
                ("ID", [f["superintelligence_vision"]["id"], f["family_daily_active_people_dec_2025"]["id"]])])
    punkt(doc, [("Okulary AI i urządzenia noszone: ", {"pogrub": True}),
                "w 2025 r. weszły okulary Meta Ray-Ban Display z wyświetlaczem w szkle, a w 2026 r. ok. "
                f"{liczba(f['rl_wearables_opex_share_2026']['wartosc'])}% kosztów Reality Labs ma trafić na urządzenia noszone.",
                ("ID", [f["meta_ray_ban_display_launch"]["id"], f["rl_wearables_opex_share_2026"]["id"]])])
    punkt(doc, [("Centra danych: ", {"pogrub": True}),
                f"spółka ma {liczba(f['data_center_locations']['wartosc'])} własnych lokalizacji centrów danych "
                "na świecie, a w wybranych miejscach dodatkowo dzierżawi centra danych.",
                ("ID", [f["data_center_locations"]["id"]])])

    # Ryzyka
    naglowek_sekcji(doc, "Ryzyka")
    for k in ("risk_advertising_dependence", "risk_ai_investments", "risk_regulation"):
        w = f[k]
        opis = w["opis"].removeprefix("Ryzyko: ")
        punkt(doc, [(w["wartosc"][0].upper() + w["wartosc"][1:] + ": ", {"pogrub": True}), opis + ".", ("ID", [w["id"]])])

    # Źródła
    naglowek_sekcji(doc, "Źródła")
    strony = {}
    for w in D.wpisy:
        if w["klasa"] == "fakt" and w["id"] in {x["id"] for x in f.values()}:
            strony.setdefault(w["zrodlo"]["strona"], []).append(w["id"])
    opis_stron = "; ".join(f"{'okładka' if isinstance(s, str) else 's. ' + str(s)}: {', '.join(v)}"
                           for s, v in sorted(strony.items(), key=lambda x: (isinstance(x[0], str), x[0])))
    p = akapit(doc, odstep_po=2)
    run(p, f"{f['revenue_2025']['zrodlo']['dokument']} (SEC EDGAR), strony i wpisy ledgera: {opis_stron}. "
           f"{f['share_price']['zrodlo']['dokument']} [{f['share_price']['id']}], pobrano {f['share_price']['data_pobrania']}. "
           f"Wycena: model {W['plik_modelu']} [wycena], założenia w ledgerze (wpisy Z). "
           "Wszystkie wpisy ledgera mają dosłowny cytat ze źródła i zostały sprawdzone dwiema bramkami.",
        ROZMIAR_DROBNY, kolor=SZARY)
    p = akapit(doc, odstep_po=0)
    run(p, ZASTRZEZENIE, ROZMIAR_DROBNY, pogrub=True, kolor=CIEMNY)
    return doc


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Użycie: python narzedzia/generuj_docx.py TICKER")
    ticker = sys.argv[1].upper()
    wymagaj_weryfikacji(ticker)
    D = Dane(ticker)
    doc = zbuduj(D)
    doc.core_properties.title = f"{ticker} - notatka dla zarządu"
    doc.core_properties.language = "pl-PL"

    katalog = KATALOG_GLOWNY / ticker / "wyniki"
    plik = katalog / f"{ticker}_notatka_zarzad.docx"
    doc.save(plik)

    with tempfile.TemporaryDirectory() as tmp:
        pdf = konwertuj(plik, "pdf", tmp)
        with pymupdf.open(pdf) as d:
            stron = len(d)
    podglad_png(plik, katalog / "podglad", "docx-strona", dpi=110)
    print(f"Notatka: {plik}; PDF ma {stron} stron(y)")
    if stron != 1:
        sys.exit(f"STOP: notatka ma {stron} strony zamiast 1. Skróć tekst (czcionka nie mniejsza niż 9 pt).")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    main()
