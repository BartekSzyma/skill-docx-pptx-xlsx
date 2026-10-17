"""Prezentacja o spółce, zbudowana z ledgera, wyniku wyceny i planu treści.

Użycie: python .claude/skills/pakiet-analityczny/scripts/generuj_pptx.py TICKER

Wejście:
  TICKER/ledger.json            - fakty, dane rynkowe, założenia, grafiki (zweryfikowane),
  TICKER/wyniki/wycena.json     - wynik modelu DCF (generuj_xlsx.py),
  TICKER/tresc.json             - plan treści: teksty z polami {ID} (format w references/tresc.md).
Wynik:
  TICKER/wyniki/TICKER_prezentacja.pptx (16:9) i podgląd TICKER/wyniki/podglad/slajd-NN.png.

Liczby wstawia skrypt z pól {ID}; przy każdym tekście dopisuje drobnym drukiem id użytych wpisów.
"""

import json
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from bramka import KATALOG_GLOWNY, wymagaj_weryfikacji
from formaty import NBSP, liczba, wartosc_z_jedn, wypelnij
from libreoffice import podglad_png

SKRYPTY = ".claude/skills/pakiet-analityczny/scripts"

# Wygląd
CIEMNY = RGBColor(0x1E, 0x2A, 0x36)
AKCENT = RGBColor(0x0F, 0x7C, 0x80)
AKCENT_JASNY = RGBColor(0x8F, 0xC9, 0xC6)
TLO_KARTY = RGBColor(0xF1, 0xF4, 0xF6)
TEKST = RGBColor(0x26, 0x32, 0x3C)
SZARY = RGBColor(0x6B, 0x77, 0x85)
CZERWONY = RGBColor(0xB0, 0x3A, 0x2E)
BIALY = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
SZER, WYS = Inches(13.333), Inches(7.5)
MARGINES = Inches(0.6)
FORMAT_LICZB = '#,##0;-#,##0'


# ---------------------------------------------------------------- dane


class Dane:
    def __init__(self, ticker):
        self.ticker = ticker
        katalog = KATALOG_GLOWNY / ticker
        self.wpisy = json.loads((katalog / "ledger.json").read_text(encoding="utf-8"))
        self.po_kluczu = {w["klucz"]: w for w in self.wpisy}
        self.po_id = {w["id"]: w for w in self.wpisy}
        plik_wyceny = katalog / "wyniki" / "wycena.json"
        if not plik_wyceny.exists():
            sys.exit(f"STOP: brak {plik_wyceny}. Najpierw uruchom: python {SKRYPTY}/generuj_xlsx.py {ticker}")
        self.wycena = json.loads(plik_wyceny.read_text(encoding="utf-8"))
        status = json.loads((katalog / "weryfikacja" / "status.json").read_text(encoding="utf-8"))
        if self.wycena.get("sha256_ledgera") != status["sha256_ledgera"]:
            sys.exit(f"STOP: wycena.json powstała z innej wersji ledgera. Uruchom ponownie: python {SKRYPTY}/generuj_xlsx.py {ticker}")
        plik_tresci = katalog / "tresc.json"
        if not plik_tresci.exists():
            sys.exit(f"STOP: brak {plik_tresci} (plan treści prezentacji i notatki, format w references/tresc.md).")
        self.tresc = json.loads(plik_tresci.read_text(encoding="utf-8"))
        self.rok = self.wycena["rok_bazowy"]

    def __getitem__(self, klucz):
        if klucz not in self.po_kluczu:
            sys.exit(f"STOP: w ledgerze {self.ticker} brakuje wpisu o kluczu '{klucz}'.")
        return self.po_kluczu[klucz]

    def wypelnij(self, tekst, dodatkowe_id=()):
        wynik, ids = wypelnij(tekst, self.po_id, self.wycena)
        for i in dodatkowe_id:
            if i != "wycena" and i not in self.po_id:
                sys.exit(f"STOP: tresc.json odwołuje się do wpisu {i}, którego nie ma w ledgerze.")
        return wynik, list(dict.fromkeys(ids + list(dodatkowe_id)))

    def spolka(self):
        return self[f"revenue_{self.rok}"]["zrodlo"]["dokument"].split(" - ")[0]


# ---------------------------------------------------------------- elementy slajdów


def tekst(slajd, x, y, w, h, tresc, rozmiar=16, kolor=TEKST, pogrub=False, wyrownanie=PP_ALIGN.LEFT,
          kotwica=MSO_ANCHOR.TOP):
    """Pole tekstowe; tresc to napis albo lista akapitów, akapit to napis albo lista fragmentów (tekst, opcje)."""
    pole = slajd.shapes.add_textbox(x, y, w, h)
    tf = pole.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = kotwica
    for i, akapit in enumerate(tresc if isinstance(tresc, list) else [tresc]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = wyrownanie
        fragmenty = akapit if isinstance(akapit, list) else [(akapit, {})]
        for t, opcje in fragmenty:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opcje.get("rozmiar", rozmiar))
            f.bold = opcje.get("pogrub", pogrub)
            f.color.rgb = opcje.get("kolor", kolor)
        pierwsze = fragmenty[0][1]
        if "odstep" in pierwsze:
            p.space_after = Pt(pierwsze["odstep"])
        if pierwsze.get("punkt"):
            ustaw_punktor(p, pierwsze.get("rozmiar", rozmiar))
    return pole


def ustaw_punktor(p, rozmiar):
    """Punktor akapitu (kwadrat w kolorze akcentu) z wcięciem wiszącym."""
    pPr = p._p.get_or_add_pPr()
    wciecie = int(Pt(rozmiar) * 1.4)
    pPr.set("marL", str(wciecie))
    pPr.set("indent", str(-wciecie))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(buClr.makeelement(qn("a:srgbClr"), {"val": str(AKCENT)}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "75000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■"}))


def id_tag(ids):
    return (f"  [{', '.join(ids)}]", {"rozmiar": 10, "kolor": SZARY, "pogrub": False}) if ids else ("", {})


def punkt(tresc_, ids, rozmiar=16, odstep=10):
    return [(tresc_, {"rozmiar": rozmiar, "punkt": True, "odstep": odstep}), id_tag(ids)]


def punkty_z_tresci(D, lista, rozmiar=16, odstep=10):
    wynik = []
    for p in lista:
        t, ids = D.wypelnij(p["tekst"], p.get("id", []))
        wynik.append(punkt(t, ids, rozmiar, odstep))
    return wynik


def prostokat(slajd, x, y, w, h, kolor):
    ksztalt = slajd.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    ksztalt.fill.solid()
    ksztalt.fill.fore_color.rgb = kolor
    ksztalt.line.fill.background()
    ksztalt.shadow.inherit = False
    ksztalt.adjustments[0] = 0.08
    return ksztalt


def tytul(slajd, napis, podtytul=None):
    tekst(slajd, MARGINES, Inches(0.45), SZER - 2 * MARGINES, Inches(0.75), napis, rozmiar=34, kolor=CIEMNY, pogrub=True)
    if podtytul:
        tekst(slajd, MARGINES, Inches(1.15), SZER - 2 * MARGINES, Inches(0.4), podtytul, rozmiar=15, kolor=SZARY)


def stopka(slajd, numer, ticker, ciemne=False):
    kolor = AKCENT_JASNY if ciemne else SZARY
    tekst(slajd, MARGINES, WYS - Inches(0.45), Inches(8), Inches(0.3),
          f"{ticker} - pakiet analityczny. Dane: ledger.json i wycena.json; [id] = wpis ledgera.", rozmiar=10, kolor=kolor)
    tekst(slajd, SZER - MARGINES - Inches(1), WYS - Inches(0.45), Inches(1), Inches(0.3), str(numer), rozmiar=10,
          kolor=kolor, wyrownanie=PP_ALIGN.RIGHT)


def styl_wykresu(wykres, legenda=True, rozmiar=12):
    wykres.font.name = FONT
    wykres.font.size = Pt(rozmiar)
    wykres.font.color.rgb = TEKST
    wykres.has_legend = legenda
    if legenda:
        wykres.legend.position = XL_LEGEND_POSITION.BOTTOM
        wykres.legend.include_in_layout = False
        wykres.legend.font.size = Pt(rozmiar)
    os_w = wykres.value_axis
    os_w.has_major_gridlines = True
    os_w.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE6)
    os_w.format.line.fill.background()
    os_w.tick_labels.font.size = Pt(rozmiar - 1)
    os_w.tick_labels.font.color.rgb = SZARY
    os_w.tick_labels.number_format = FORMAT_LICZB
    os_w.tick_labels.number_format_is_linked = False
    wykres.category_axis.format.line.color.rgb = RGBColor(0xBF, 0xC6, 0xCC)
    wykres.category_axis.tick_labels.font.size = Pt(rozmiar)


def etykiety(seria, format_liczb=FORMAT_LICZB, rozmiar=11):
    seria.data_labels.show_value = True
    seria.data_labels.number_format = format_liczb
    seria.data_labels.number_format_is_linked = False
    seria.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    seria.data_labels.font.size = Pt(rozmiar)
    seria.data_labels.font.color.rgb = TEKST


def tylko_dodatnie(wartosci, gdzie):
    # LibreOffice rysuje ujemne słupki z python-pptx jako dodatnie, więc podgląd byłby mylący.
    if any(v < 0 for v in wartosci):
        sys.exit(f"STOP: {gdzie}: wykres słupkowy z wartością ujemną. LibreOffice pokazuje ją jako dodatnią; "
                 "pokaż na wykresie wartości dodatnie (np. przychody), a stratę opisz w punktach.")


# ---------------------------------------------------------------- slajdy


def slajd_tytulowy(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = CIEMNY
    rev, kurs = D[f"revenue_{D.rok}"], D["share_price"]
    tekst(s, MARGINES, Inches(2.0), Inches(12), Inches(1.0), D.spolka(), rozmiar=44, kolor=BIALY, pogrub=True)
    tekst(s, MARGINES, Inches(3.0), Inches(11), Inches(0.6), f"Pakiet analityczny  |  {D.ticker}", rozmiar=26, kolor=AKCENT_JASNY)
    tekst(s, MARGINES, Inches(4.3), Inches(12), Inches(1.2), [
        [(f"Dane finansowe: {rev['zrodlo']['dokument']}", {"rozmiar": 16, "kolor": BIALY}), id_tag([rev["id"]])],
        [(f"Dane rynkowe: kurs z dnia {kurs['okres']}, pobrane {kurs['data_pobrania']}", {"rozmiar": 16, "kolor": BIALY}),
         id_tag([kurs["id"]])],
        [(f"Wycena DCF z dnia {D.wycena['data_wyceny']}", {"rozmiar": 16, "kolor": BIALY}), id_tag(["wycena"])],
    ])
    stopka(s, nr, D.ticker, ciemne=True)


def slajd_liczby(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    lata = [str(D.rok - 2), str(D.rok - 1), str(D.rok)]
    rev = [D[f"revenue_{r}"] for r in lata]
    opi = [D[f"operating_income_{r}"] for r in lata]
    tytul(s, "Spółka w liczbach", f"Przychody i wynik operacyjny w latach {lata[0]}-{lata[-1]}, mln USD")
    tylko_dodatnie([w["wartosc"] for w in rev], "slajd 'Spółka w liczbach' (przychody)")
    # Strata operacyjna nie trafia na wykres (LibreOffice rysuje ujemne słupki jako dodatnie); wtedy podajemy ją w podpisie.
    strata = any(w["wartosc"] < 0 for w in opi)
    cd = CategoryChartData()
    cd.categories = lata
    cd.add_series("Przychody", [w["wartosc"] for w in rev])
    if not strata:
        cd.add_series("Zysk operacyjny", [w["wartosc"] for w in opi])
    wykres = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGINES, Inches(1.75), Inches(7.6), Inches(4.45), cd).chart
    styl_wykresu(wykres, legenda=not strata)
    wykres.plots[0].gap_width = 60
    wykres.plots[0].overlap = -10
    for seria, kolor in zip(wykres.series, (AKCENT, CIEMNY)):
        seria.format.fill.solid()
        seria.format.fill.fore_color.rgb = kolor
        etykiety(seria)
    strony = sorted({str(w["zrodlo"]["strona"]) for w in rev + opi})
    podpis = f"Źródło: rachunek zysków i strat, raport 10-K, s. {', '.join(strony)}"
    if strata:
        podpis = ("Wynik operacyjny: " + "; ".join(f"{r}: {wartosc_z_jedn(w)}" for r, w in zip(lata, opi)) + ". " + podpis)
    tekst(s, MARGINES, Inches(6.3), Inches(7.6), Inches(0.45),
          [[(podpis, {"rozmiar": 10, "kolor": SZARY}), id_tag([w["id"] for w in rev + opi])]])

    karty = D.tresc["prezentacja"]["karty"][:3]
    for i, karta in enumerate(karty):
        y = Inches(1.75 + 1.6 * i)
        wartosc, ids_w = D.wypelnij(karta["wartosc"])
        opis, ids_o = D.wypelnij(karta["opis"], karta.get("id", []))
        prostokat(s, Inches(8.7), y, Inches(4.0), Inches(1.4), TLO_KARTY)
        tekst(s, Inches(8.95), y + Inches(0.18), Inches(3.5), Inches(0.65), wartosc, rozmiar=30, kolor=AKCENT, pogrub=True)
        tekst(s, Inches(8.95), y + Inches(0.85), Inches(3.5), Inches(0.5),
              [[(opis, {"rozmiar": 14}), id_tag(list(dict.fromkeys(ids_w + ids_o)))]])
    stopka(s, nr, D.ticker)


def slajd_dodatkowy(prs, D, nr, cfg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, cfg["tytul"], cfg.get("podtytul"))
    wykres_cfg = cfg.get("wykres")
    szer_tekstu = Inches(6.0) if wykres_cfg else Inches(12.1)
    tekst(s, MARGINES, Inches(1.85), szer_tekstu, Inches(4.5), punkty_z_tresci(D, cfg["punkty"]))
    if wykres_cfg:
        wpisy = [D.po_id[i] for i in wykres_cfg["wartosci"]]
        wartosci = [w["wartosc"] for w in wpisy]
        tylko_dodatnie(wartosci, f"slajd '{cfg['tytul']}'")
        cd = CategoryChartData()
        cd.categories = wykres_cfg["kategorie"]
        cd.add_series(wykres_cfg["tytul"], wartosci)
        wykres = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.2), Inches(1.75), Inches(5.5), Inches(4.65), cd).chart
        styl_wykresu(wykres, legenda=False)
        wykres.has_title = True
        wykres.chart_title.text_frame.text = wykres_cfg["tytul"]
        wykres.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
        wykres.chart_title.text_frame.paragraphs[0].runs[0].font.bold = False
        seria = wykres.series[0]
        seria.format.fill.solid()
        seria.format.fill.fore_color.rgb = AKCENT
        etykiety(seria)
        tekst(s, Inches(7.2), Inches(6.45), Inches(5.5), Inches(0.3),
              [[(f"Źródło: {wykres_cfg.get('zrodlo', 'raport 10-K')}", {"rozmiar": 10, "kolor": SZARY}),
                id_tag([w["id"] for w in wpisy])]])
    stopka(s, nr, D.ticker)


def slajd_projekt(prs, D, nr, cfg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    g = D.po_id.get(cfg["grafika"])
    if g is None or g["klasa"] != "grafika":
        sys.exit(f"STOP: tresc.json wskazuje grafikę {cfg['grafika']}, której nie ma w ledgerze.")
    tytul(s, cfg["tytul"], cfg.get("podtytul"))
    plik = KATALOG_GLOWNY / g["plik"]
    with Image.open(plik) as obraz:
        szer, wys = obraz.size
    skala = min(Inches(6.4) / szer, Inches(4.3) / wys)
    w, h = int(szer * skala), int(wys * skala)
    s.shapes.add_picture(str(plik), MARGINES + (Inches(6.4) - w) // 2, Inches(1.75), w, h)
    tekst(s, MARGINES, Inches(1.75) + h + Inches(0.08), Inches(6.4), Inches(0.5), [
        [(f"{g['opis']}. Fot. {g['autor']}, licencja {g['licencja']}, Wikimedia Commons", {"rozmiar": 10, "kolor": SZARY}),
         id_tag([g["id"]])]])
    prostokat(s, Inches(7.35), Inches(1.75), Inches(5.4), Inches(4.65), TLO_KARTY)
    tekst(s, Inches(7.65), Inches(2.05), Inches(4.8), Inches(4.1), punkty_z_tresci(D, cfg["punkty"], rozmiar=17, odstep=14))
    stopka(s, nr, D.ticker)


def slajd_wycena(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W = D.wycena
    p0, pn = D.rok + 1, D.rok + 5
    tytul(s, "Wycena DCF: wartość wobec kursu", f"Wynik modelu z pliku {D.ticker}_DCF.xlsx; zależy od przyjętych założeń")
    cd = CategoryChartData()
    cd.categories = ["Wartość z DCF", "Kurs rynkowy"]
    cd.add_series("USD za akcję", [W["wartosc_jednej_akcji_usd"], W["kurs_usd"]])
    tylko_dodatnie([W["wartosc_jednej_akcji_usd"]], "slajd wyceny (ujemna wartość kapitału)")
    wykres = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGINES, Inches(1.75), Inches(6.2), Inches(2.6), cd).chart
    styl_wykresu(wykres, legenda=False)
    wykres.value_axis.visible = False
    wykres.value_axis.has_major_gridlines = False
    wykres.value_axis.minimum_scale = 0
    wykres.category_axis.reverse_order = True
    wykres.plots[0].gap_width = 45
    seria = wykres.series[0]
    for i, kolor in enumerate((AKCENT, CIEMNY)):
        seria.points[i].format.fill.solid()
        seria.points[i].format.fill.fore_color.rgb = kolor
    etykiety(seria, '#,##0.00" USD"', rozmiar=14)
    kurs = D["share_price"]
    tekst(s, MARGINES, Inches(4.45), Inches(6.2), Inches(1.1), [
        [(f"{liczba(W['roznica_proc'], 1)}%", {"rozmiar": 36, "kolor": CZERWONY if W["roznica_proc"] < 0 else AKCENT, "pogrub": True}),
         ("  różnica wartości z DCF wobec kursu", {"rozmiar": 16}), id_tag(["wycena", kurs["id"]])]])
    tekst(s, MARGINES, Inches(5.4), Inches(6.2), Inches(1.0), [
        [(f"Wartość firmy {liczba(W['wartosc_firmy_mln_usd'])}{NBSP}mln USD, kapitał własny "
          f"{liczba(W['wartosc_kapitalu_mln_usd'])}{NBSP}mln USD; {liczba(W['udzial_wartosci_rezydualnej_proc'])}% wartości "
          f"firmy pochodzi z okresu po {pn} r.", {"rozmiar": 14}), id_tag(["wycena"])]])

    x, w = Inches(7.25), Inches(5.5)
    prostokat(s, x, Inches(1.75), w, Inches(4.65), TLO_KARTY)
    tekst(s, x + Inches(0.3), Inches(1.95), w - Inches(0.6), Inches(0.75), [
        [("To są ZAŁOŻENIA, nie fakty", {"rozmiar": 18, "pogrub": True, "kolor": CZERWONY})],
        [("Przyjęte przez nas, z uzasadnieniem w ledgerze; zmiana założeń zmienia wynik.", {"rozmiar": 12, "kolor": SZARY})]])
    z = {k: D[k] for k in (f"revenue_growth_{p0}", f"revenue_growth_{pn}", "operating_margin_terminal",
                           f"capex_pct_revenue_{p0}", "beta", "equity_risk_premium", "terminal_growth")}
    rf = D["us_treasury_10y_yield"]
    pozycje = [
        (f"Wzrost przychodów: {wartosc_z_jedn(z[f'revenue_growth_{p0}'])} w {p0} r., "
         f"{wartosc_z_jedn(z[f'revenue_growth_{pn}'])} w {pn} r.", [z[f"revenue_growth_{p0}"]["id"], z[f"revenue_growth_{pn}"]["id"]]),
        (f"Marża operacyjna po {pn} r.: {wartosc_z_jedn(z['operating_margin_terminal'])}", [z["operating_margin_terminal"]["id"]]),
        (f"Nakłady inwestycyjne w {p0} r.: {wartosc_z_jedn(z[f'capex_pct_revenue_{p0}'])} przychodów",
         [z[f"capex_pct_revenue_{p0}"]["id"]]),
        (f"WACC {liczba(W['wacc_proc'], 1)}% = rentowność obligacji {liczba(rf['wartosc'], 2)}% + beta "
         f"{liczba(z['beta']['wartosc'], 1)} × premia {liczba(z['equity_risk_premium']['wartosc'], 1)}%",
         ["wycena", rf["id"], z["beta"]["id"], z["equity_risk_premium"]["id"]]),
        (f"Wzrost po {pn} r.: {liczba(z['terminal_growth']['wartosc'], 1)}% rocznie", [z["terminal_growth"]["id"]]),
    ]
    tekst(s, x + Inches(0.3), Inches(2.95), w - Inches(0.6), Inches(3.3), [punkt(t, ids, 15, 9) for t, ids in pozycje])
    stopka(s, nr, D.ticker)


def slajd_ryzyka(prs, D, nr):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tytul(s, "Ryzyka i źródła")
    ryzyka = [w for w in D.wpisy if w["klasa"] == "fakt" and w["klucz"].startswith("risk_")][:3]
    for i, w in enumerate(ryzyka):
        y = Inches(1.4 + 1.6 * i)
        opis = w["opis"].removeprefix("Ryzyko: ")
        prostokat(s, MARGINES, y, Inches(6.6), Inches(1.45), TLO_KARTY)
        tekst(s, MARGINES + Inches(0.25), y + Inches(0.15), Inches(6.1), Inches(1.2), [
            [(str(w["wartosc"])[0].upper() + str(w["wartosc"])[1:], {"rozmiar": 16, "pogrub": True, "kolor": CZERWONY})],
            [(opis[0].upper() + opis[1:], {"rozmiar": 13}), id_tag([w["id"]]),
             (f"  10-K, s. {w['zrodlo']['strona']}", {"rozmiar": 10, "kolor": SZARY})]])

    strony = sorted({w["zrodlo"]["strona"] for w in D.wpisy if w["klasa"] == "fakt"},
                    key=lambda s_: (not isinstance(s_, int), str(s_) if not isinstance(s_, int) else s_))
    strony_tekst = ", ".join("okładka" if str(s_).startswith("okładka") else str(s_) for s_ in strony)
    dok = D[f"revenue_{D.rok}"]["zrodlo"]
    zrodla = [[(dok["dokument"], {"rozmiar": 14, "pogrub": True, "odstep": 2})],
              [(f"SEC EDGAR, strony: {strony_tekst}", {"rozmiar": 12, "odstep": 10})]]
    for w in (w for w in D.wpisy if w["klasa"] == "dana_rynkowa"):
        zrodla += [[(w["zrodlo"]["dokument"], {"rozmiar": 14, "pogrub": True, "odstep": 2})],
                   [(f"{w['opis']}; stan na {w['okres']}, pobrano {w['data_pobrania']}", {"rozmiar": 12, "odstep": 10}),
                    id_tag([w["id"]])]]
    graf = [w for w in D.wpisy if w["klasa"] == "grafika"]
    if graf:
        zrodla += [[("Wikimedia Commons - zdjęcia", {"rozmiar": 14, "pogrub": True, "odstep": 2})],
                   [("; ".join(f"{g['autor']} ({g['licencja']})" for g in graf), {"rozmiar": 12, "odstep": 10}),
                    id_tag([g["id"] for g in graf])]]
    zrodla += [[("Ledger i weryfikacja", {"rozmiar": 14, "pogrub": True, "odstep": 2})],
               [(f"{D.ticker}/ledger.json: każdy wpis ze źródłem i cytatem, sprawdzony dwiema bramkami "
                 "(skrypt i niezależny model)", {"rozmiar": 12})]]
    tekst(s, Inches(7.6), Inches(1.4), Inches(5.1), Inches(0.4), "Źródła", rozmiar=18, kolor=CIEMNY, pogrub=True)
    tekst(s, Inches(7.6), Inches(1.9), Inches(5.1), Inches(4.6), zrodla)
    stopka(s, nr, D.ticker)


# ---------------------------------------------------------------- main


def main():
    if len(sys.argv) != 2:
        raise SystemExit(f"Użycie: python {SKRYPTY}/generuj_pptx.py TICKER")
    ticker = sys.argv[1].upper()
    wymagaj_weryfikacji(ticker)
    D = Dane(ticker)
    cfg = D.tresc["prezentacja"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = SZER, WYS
    prs.core_properties.title = f"{ticker} - pakiet analityczny"
    prs.core_properties.language = "pl-PL"

    nr = 1
    slajd_tytulowy(prs, D, nr)
    nr += 1
    slajd_liczby(prs, D, nr)
    if cfg.get("slajd_dodatkowy"):
        nr += 1
        slajd_dodatkowy(prs, D, nr, cfg["slajd_dodatkowy"])
    for projekt in cfg.get("projekty", []):
        nr += 1
        slajd_projekt(prs, D, nr, projekt)
    nr += 1
    slajd_wycena(prs, D, nr)
    nr += 1
    slajd_ryzyka(prs, D, nr)

    katalog = KATALOG_GLOWNY / ticker / "wyniki"
    plik = katalog / f"{ticker}_prezentacja.pptx"
    prs.save(plik)
    liczba_png = podglad_png(plik, katalog / "podglad", "slajd", dpi=80)
    print(f"Prezentacja: {plik} ({len(prs.slides)} slajdów), podgląd: {liczba_png} plików PNG w {katalog / 'podglad'}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    main()
