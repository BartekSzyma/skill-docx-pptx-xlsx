"""Kontrola liczb w dokumentach: każda liczba w pptx i docx musi mieć pokrycie w ledgerze albo w wycenie.

Użycie: python .claude/skills/pakiet-analityczny/scripts/sprawdz_liczby.py TICKER

Sprawdza TICKER/wyniki/*.pptx (tekst, tabele, dane wykresów) i *.docx (akapity, tabele).
Liczba ma pokrycie, jeśli po zaokrągleniu do liczby miejsc po przecinku z tekstu równa się:
  - wartości wpisu ledgera (także bez znaku; kwoty w mln/mld USD także po zamianie mln <-> mld),
  - wartości z wyniki/wycena.json,
  - numerowi strony źródła albo liczbie z pola tekstowego wpisu (np. "CC BY 4.0" w licencji).
Pomija: lata (1990-2100 bez separatora tysięcy), daty RRRR-MM-DD, id wpisów (F-01, R-02...),
nazwę formularza 10-K i numery slajdów w stopce.
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation

from bramka import KATALOG_GLOWNY

LICZBA = re.compile(r"(?<![\w,.])(-?\d{1,3}(?:[   ]\d{3})+(?:,\d+)?|-?\d+(?:[.,]\d+)?)(?![\w])")
DO_POMINIECIA = [
    re.compile(r"\b[FRZG]-\d{2}\b"),               # id wpisów
    re.compile(r"\b\d{4}-\d{2}-\d{2}\b"),           # daty
    re.compile(r"\b\d{4}-\d{2}\b"),                 # rok-miesiąc
    re.compile(r"\b10-K\b"),                        # nazwa formularza
]
LICZBA_W_TEKSCIE = re.compile(r"\d+(?:[.,]\d+)?")


def zamien(tekst):
    """'200 966' -> 200966.0, '81,59' -> 81.59, '4.0' -> 4.0; zwraca też liczbę miejsc po przecinku."""
    t = re.sub(r"[   ]", "", tekst).replace(",", ".")
    miejsca = len(t.split(".")[1]) if "." in t else 0
    return float(t), miejsca


def pokrycie(ticker):
    """Lista (wartość, opis źródła) dopuszczalnych liczb."""
    katalog = KATALOG_GLOWNY / ticker
    wpisy = json.loads((katalog / "ledger.json").read_text(encoding="utf-8"))
    wycena = json.loads((katalog / "wyniki" / "wycena.json").read_text(encoding="utf-8"))
    kandydaci = []
    for w in wpisy:
        v = w.get("wartosc")
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            kandydaci.append((v, w["id"]))
            # Zamiana mln <-> mld tylko dla kwot w tych jednostkach (inaczej np. kurs/1000 "pokrywałby" liczbę 1).
            if w.get("jednostka") == "mln USD":
                kandydaci.append((v / 1000, f"{w['id']} w mld"))
            elif w.get("jednostka") == "mld USD":
                kandydaci.append((v * 1000, f"{w['id']} w mln"))
        strona = w.get("zrodlo", {}).get("strona")
        if isinstance(strona, int):
            kandydaci.append((strona, f"{w['id']} (strona źródła)"))
        for pole in ("opis", "licencja", "autor", "wartosc"):
            if isinstance(w.get(pole), str):
                for t in LICZBA_W_TEKSCIE.findall(w[pole]):
                    kandydaci.append((float(t.replace(",", ".")), f"{w['id']} (pole {pole})"))
    for klucz, v in wycena.items():
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            kandydaci.append((v, f"wycena.{klucz}"))
    return kandydaci


def znajdz_pokrycie(wartosc, miejsca, kandydaci):
    # Porównujemy wartości bezwzględne: nakłady zapisane w ledgerze jako wypływ (-69691) w tekście są dodatnie.
    for v, opis in kandydaci:
        if abs(round(abs(v), miejsca) - abs(wartosc)) < 1e-9:
            return opis
    return None


def liczby_w_tekscie(tekst):
    for wzor in DO_POMINIECIA:
        tekst = wzor.sub(" ", tekst)
    for m in LICZBA.finditer(tekst):
        surowa = m.group(1)
        wartosc, miejsca = zamien(surowa)
        zgrupowana = bool(re.search(r"\d[   ]\d", surowa))
        if not zgrupowana and miejsca == 0 and 1990 <= abs(wartosc) <= 2100:
            continue  # rok
        yield surowa, wartosc, miejsca


def teksty_pptx(plik):
    prs = Presentation(plik)
    for nr, slajd in enumerate(prs.slides, start=1):
        def ksztalty(kolekcja):
            for k in kolekcja:
                if k.shape_type == 6:  # grupa
                    yield from ksztalty(k.shapes)
                else:
                    yield k
        for k in ksztalty(slajd.shapes):
            if k.has_text_frame:
                tekst = k.text_frame.text
                if tekst.strip().isdigit() and int(tekst.strip()) == nr:
                    continue  # numer slajdu w stopce
                yield f"slajd {nr}, pole '{k.name}'", tekst
            if getattr(k, "has_table", False) and k.has_table:
                for r, wiersz in enumerate(k.table.rows, start=1):
                    for c, komorka in enumerate(wiersz.cells, start=1):
                        yield f"slajd {nr}, tabela '{k.name}' w{r} k{c}", komorka.text
            if getattr(k, "has_chart", False) and k.has_chart:
                wykres = k.chart
                if wykres.has_title:
                    yield f"slajd {nr}, tytuł wykresu", wykres.chart_title.text_frame.text
                for wykres_plot in wykres.plots:
                    yield f"slajd {nr}, kategorie wykresu", " ; ".join(str(c) for c in wykres_plot.categories)
                    for seria in wykres_plot.series:
                        for v in seria.values:
                            # Wartości wykresu sprawdzamy w zapisie z kropką (tak są zapisane w pliku).
                            yield f"slajd {nr}, wykres '{seria.name}'", repr(float(v)).removesuffix(".0")


def teksty_docx(plik):
    doc = Document(plik)
    for i, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            yield f"akapit {i}", p.text
    for t, tabela in enumerate(doc.tables, start=1):
        for r, wiersz in enumerate(tabela.rows, start=1):
            for c, komorka in enumerate(wiersz.cells, start=1):
                yield f"tabela {t} w{r} k{c}", komorka.text


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Użycie: python .claude/skills/pakiet-analityczny/scripts/sprawdz_liczby.py TICKER")
    ticker = sys.argv[1].upper()
    katalog = KATALOG_GLOWNY / ticker / "wyniki"
    kandydaci = pokrycie(ticker)
    pliki = sorted(katalog.glob("*.pptx")) + sorted(katalog.glob("*.docx"))
    if not pliki:
        sys.exit(f"Brak plików pptx/docx w {katalog}")

    bledy, sprawdzone = [], 0
    for plik in pliki:
        zrodlo = teksty_pptx(plik) if plik.suffix == ".pptx" else teksty_docx(plik)
        w_pliku = 0
        for miejsce, tekst in zrodlo:
            for surowa, wartosc, miejsca in liczby_w_tekscie(tekst):
                sprawdzone += 1
                w_pliku += 1
                if znajdz_pokrycie(wartosc, miejsca, kandydaci) is None:
                    fragment = " ".join(tekst.split())
                    poz = fragment.find(surowa.replace(" ", " "))
                    kontekst = fragment[max(0, poz - 40):poz + len(surowa) + 40] if poz >= 0 else fragment[:80]
                    bledy.append(f"{plik.name}, {miejsce}: liczba '{surowa}' bez pokrycia - ...{kontekst}...")
        print(f"{plik.name}: sprawdzono {w_pliku} liczb")

    if bledy:
        print(f"\nBŁĘDY ({len(bledy)}):")
        for b in bledy:
            print("  " + b)
        sys.exit(1)
    print(f"\nKontrola liczb: OK ({sprawdzone} liczb w {len(pliki)} plikach ma pokrycie w ledgerze albo wycenie)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
